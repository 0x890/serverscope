from __future__ import annotations

from pathlib import Path

from serverscope.report.base import ReportContext, Reporter


class ExcelReporter(Reporter):
    name = "excel"
    extension = "xlsx"

    def render(self, context: ReportContext, output_dir: Path) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError as exc:
            raise RuntimeError("openpyxl required: pip install serverscope[excel]") from exc

        wb = Workbook()
        ws = wb.active
        ws.title = "Inventory"
        header = ["Server", "Platform", "Classification", "Health", "Risk",
                  "Idle %", "Abandonment %", "Perf", "Sec", "Avail", "Recommendation"]
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2563EB")
        for r in context.results:
            ws.append([
                r.hostname, r.platform.value, r.classification.value,
                r.scores.health, r.scores.risk, r.scores.idle_probability,
                r.scores.abandonment_probability, r.performance_grade,
                r.security_grade, r.availability_grade, r.recommendation,
            ])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        summary = wb.create_sheet("Summary")
        summary.append(["Generated", context.generated_at.isoformat()])
        summary.append(["Total servers", context.total])
        summary.append(["Average health", context.avg_health])
        for cls, count in context.classification_counts.most_common():
            summary.append([cls, count])

        findings = wb.create_sheet("Critical Findings")
        findings.append(["Server", "Finding"])
        for host, desc in context.critical_anomalies:
            findings.append([host, desc])

        path = self.output_path(output_dir)
        wb.save(path)
        return path
